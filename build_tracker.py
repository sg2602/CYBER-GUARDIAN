"""Build CyberGuardian_Project_Tracker.xlsx from the real result files.

    python build_tracker.py
Sheets: Task Tracker, Test Cases, ML Results, Bugs, Literature, Novelty Evidence.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
R = ROOT / "results"
OUT = ROOT / "CyberGuardian_Project_Tracker.xlsx"
ME, MATE, BOTH = "Me", "Teammate", "Both"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {"Done": "C8E6C9", "Testing": "FFF9C4", "In Progress": "BBDEFB", "Blocked": "FFCDD2", "Not Started": "EEEEEE"}

# (module, task, owner, day, priority, branch)
TASKS = [
    ("Setup", "Finalize modules + architecture diagram", BOTH, 1, "MUST", "main"),
    ("Setup", "Repo + folder structure + README skeleton", BOTH, 1, "MUST", "main"),
    ("Setup", "Define ModuleResult / Finding contract (utils/common.py)", ME, 1, "MUST", "main"),
    ("Setup", "Define finding IDs, chain rules + ATT&CK mapping, risk formula", ME, 1, "MUST", "main"),
    ("Setup", "Define DB schema and UI page list", MATE, 1, "MUST", "main"),
    ("Password", "Length / class / entropy checks", ME, 2, "MUST", "password-module"),
    ("Password", "Common-password list (SecLists 10k) + dictionary + leetspeak", ME, 2, "MUST", "password-module"),
    ("Password", "Sequence / repeat / name+digits / year patterns, reuse flag", ME, 2, "MUST", "password-module"),
    ("Device", "Questionnaire, weights, 'Not sure' = half risk", MATE, 2, "MUST", "dashboard"),
    ("URL", "URL parsing, registered domain, feature extraction", ME, 3, "MUST", "url-module"),
    ("URL", "HTTPS / IP / TLD / @ / shortener / punycode / homograph rules", ME, 3, "MUST", "url-module"),
    ("ML-URL", "URL dataset collection + preprocessing", MATE, 3, "MUST", "url-module"),
    ("ML-URL", "Train RF + XGBoost, evaluate, serialize", MATE, 4, "MUST", "url-module"),
    ("URL", "Rule + ML fusion, SHAP, disagreement flag", ME, 4, "SHOULD", "url-module"),
    ("Email", "Header analysis (Reply-To, spoofing, SPF/DKIM/DMARC)", ME, 5, "MUST", "email-module"),
    ("Email", "Content rules + India-localized lexicon (C5)", ME, 5, "NICE", "email-module"),
    ("ML-Email", "Email dataset + TF-IDF + LR / RF training", MATE, 5, "MUST", "email-module"),
    ("Email", "Header + Content + URL + ML -> Email risk; test samples", BOTH, 6, "MUST", "email-module"),
    ("Permissions", "Permission -> severity table, category baseline, manifest parser", MATE, 7, "MUST", "dashboard"),
    ("Risk", "Risk engine: renormalised weights, coverage, baseline_score", ME, 8, "MUST", "risk-engine"),
    ("Correlation", "Threat Correlation Engine: chains, escalation, critical floor", ME, 8, "MUST", "correlation-engine"),
    ("Risk", "Score breakdown (points lost) + reconciliation test", ME, 8, "SHOULD", "risk-engine"),
    ("Recommendations", "Finding -> action -> priority mapping", ME, 9, "MUST", "risk-engine"),
    ("What-If", "Counterfactual simulator + ranking by gain/effort", ME, 9, "SHOULD", "risk-engine"),
    ("Dashboard", "Streamlit pages: scanner, dashboard, what-if, history, report, about", MATE, 10, "MUST", "dashboard"),
    ("Dashboard", "Attack-chain panel, waterfall, UNCERTAIN badges, what-if toggles", MATE, 10, "SHOULD", "dashboard"),
    ("Database", "SQLite persistence (IDs and scores only)", MATE, 11, "MUST", "dashboard"),
    ("History", "History page + trend chart + improvement delta", MATE, 11, "SHOULD", "dashboard"),
    ("Report", "HTML report + PDF export", MATE, 11, "SHOULD", "dashboard"),
    ("Integration", "End-to-end pipeline freeze (no new features)", BOTH, 12, "MUST", "main"),
    ("Testing", "pytest unit tests + module test suites", BOTH, 13, "MUST", "main"),
    ("Evaluation", "Scenario study (C1)", ME, 13, "MUST", "main"),
    ("Evaluation", "Adversarial URL robustness test", MATE, 13, "NICE", "main"),
    ("Evaluation", "Rule-ML disagreement statistics + India test set", ME, 13, "NICE", "main"),
    ("Final", "Bug fixes, README, screenshots, limitations, future work", BOTH, 14, "MUST", "main"),
    ("Final", "Demo rehearsal, PPT, viva prep (VIVA_PREP.md)", BOTH, 14, "MUST", "main"),
]
BUILT_STATUS = "Done"  # the code for every task above exists and is tested; PPT/rehearsal are still to do
NOT_BUILT = {"Demo rehearsal, PPT, viva prep (VIVA_PREP.md)": "In Progress",
             "Bug fixes, README, screenshots, limitations, future work": "Testing"}

BUGS = [
    ("BUG-01", "ML-URL", "Model trained on PhiUSIIL alone learned 'has a path => phishing' / 'https => legit' (99% test acc, but wikipedia/github deep links scored CRITICAL)", "High", MATE, "Fixed - mixed benign deep-link sources, randomised scheme for scheme-less benign URLs"),
    ("BUG-02", "ML-URL", "Trailing '/' on homepages (https://site/) treated as a path - PhiUSIIL artefact made iitb.ac.in/ look like phishing", "High", MATE, "Fixed - trailing slash normalised in feature extraction; model retrained"),
    ("BUG-03", "ML-URL", "Random Forest model file was 122 MB (too big to ship)", "Medium", MATE, "Fixed - limited depth/trees + joblib compression (13 MB)"),
    ("BUG-04", "Risk", "Score breakdown did not reconcile exactly (module scores rounded to 6 dp)", "Medium", ME, "Fixed - no rounding inside the engine; reconciliation unit test added"),
    ("BUG-05", "Password", "Samriddhi123 scored CRITICAL: dictionary-word and name+digits findings double counted", "Medium", ME, "Fixed - word inside Name+digits pattern no longer counted twice"),
    ("BUG-06", "Fusion", "Fixed 60/40 rule/ML blend let a confident ML phishing verdict score only 60 (MODERATE) when rules saw nothing", "High", ME, "Fixed - fusion method chosen on out-of-fold training data (ml/select_fusion.py)"),
    ("BUG-07", "Correlation", "HIGH-severity chains only capped a module tier, so the overall score stayed MODERATE (scenario study v1: 5/9 dangerous profiles mislabelled)", "High", ME, "Fixed - HIGH-chain cap at 59 (ablation kept in results)"),
    ("BUG-08", "UI", "Demo buttons set widget state after widget creation (StreamlitAPIException)", "Low", MATE, "Fixed - on_click callbacks"),
    ("BUG-09", "UI", "Home page markdown indented inside function rendered as code block", "Low", MATE, "Fixed"),
    ("BUG-10", "Password", "Reused-but-strong password still LOW risk (reuse weight 20)", "Medium", ME, "Fixed - reuse weight 25 (-> MODERATE)"),
    ("BUG-11", "ML-URL", "Legit deep links on non-allow-listed domains (stackoverflow, bbc) still scored HIGH/CRITICAL by ML", "Medium", MATE, "Open - shown as UNCERTAIN; needs better benign data"),
    ("BUG-12", "ML-Email", "Short SMS-style legit messages (OTP, delivery) get high phishing probability (dataset shift)", "Medium", MATE, "Open - documented; India rules + UNCERTAIN flag mitigate"),
    ("BUG-13", "Email", "UNCERTAIN flag fires on ~40% of URLs/emails because rules have low recall on public corpora", "Medium", ME, "Open - future work: calibrate rule scores"),
    ("BUG-14", "Correlation", "No chain for reused/weak password + no MFA without an observed lure (P06) or for 'digital arrest' alone (P13)", "Low", ME, "Open - documented limitation"),
]

LIT = [
    ("Ma, Saul, Savage, Voelker - Beyond Blacklists: Learning to Detect Malicious Web Sites from Suspicious URLs (KDD)", 2009, "Malicious URL detection", "Lexical + host features, online learning", "Own crawl", "Lexical URL features idea", "Confident"),
    ("Garera, Provos, Chew, Rubin - A Framework for Detection and Measurement of Phishing Attacks (WORM)", 2007, "Phishing URL detection", "URL structure features + logistic regression", "Google toolbar data", "Obfuscation feature families (@, IP, brand in path)", "Confident"),
    ("Sahingoz, Buber, Demir, Diri - Machine learning based phishing detection from URLs (Expert Systems with Applications)", 2019, "Phishing URL detection", "NLP features + RF", "Own dataset", "Comparison of classical ML on URL features", "Confident"),
    ("Prasad & Chandra - PhiUSIIL: A diverse security profile empowered phishing URL detection framework based on similarity index and incremental learning (Computers & Security)", 2024, "Phishing URL detection", "Similarity index + incremental learning", "PhiUSIIL (UCI 967)", "Dataset used for training", "Confident"),
    ("Mamun, Rathore, Lashkari, Stakhanova, Ghorbani - Detecting Malicious URLs Using Lexical Analysis (NSS)", 2016, "Malicious URL detection", "Lexical features", "ISCX-URL-2016", "Benign URL list source", "Verify venue details"),
    ("Fette, Sadeh, Tomasic - Learning to Detect Phishing Emails (WWW)", 2007, "Phishing email detection", "PILFER: engineered features + RF", "SpamAssassin + phishing corpus", "Header/URL engineered email features", "Confident"),
    ("Chandrasekaran, Narayanan, Upadhyaya - Phishing email detection based on structural properties (NYS Cyber Security Symposium)", 2006, "Phishing email detection", "Structural features + SVM", "Own", "Structural cues", "Verify"),
    ("Felt, Chin, Hanna, Song, Wagner - Android Permissions Demystified (ACM CCS)", 2011, "Android permission risk", "Static analysis of over-privilege", "940 apps", "Over-privilege / purpose mismatch idea", "Confident"),
    ("Enck, Ongtang, McDaniel - On Lightweight Mobile Phone Application Certification (ACM CCS)", 2009, "Android permission risk", "Kirin: rules over dangerous permission combinations", "Android apps", "Dangerous-combination rules (e.g. SMS + ...)", "Confident"),
    ("Sarma, Li, Gates, Potharaju, Nita-Rotaru, Molloy - Android Permissions: A Perspective Combining Risks and Benefits (SACMAT)", 2012, "Android permission risk", "Category-relative rare-permission risk", "Google Play apps", "Category baseline for expected permissions", "Confident"),
    ("Lundberg & Lee - A Unified Approach to Interpreting Model Predictions (NeurIPS)", 2017, "Explainability", "SHAP values", "-", "SHAP for URL/email models", "Confident"),
    ("Lundberg et al. - From local explanations to global understanding with explainable AI for trees (Nature Machine Intelligence)", 2020, "Explainability", "TreeExplainer", "-", "TreeExplainer for RF/XGBoost", "Confident"),
    ("Strom et al. - MITRE ATT&CK: Design and Philosophy (MITRE technical report)", 2018, "Threat modelling", "ATT&CK knowledge base", "-", "Technique IDs for findings and chains", "Confident"),
    ("Sheyner, Haines, Jha, Lippmann, Wing - Automated Generation and Analysis of Attack Graphs (IEEE S&P)", 2002, "Attack graphs / risk scoring", "Model checking to build attack graphs", "Network models", "Idea of chaining weaknesses (C1)", "Confident"),
    ("Ou, Govindavajhala, Appel - MulVAL: A Logic-based Network Security Analyzer (USENIX Security)", 2005, "Attack graphs / risk scoring", "Datalog rules over vulnerabilities", "Network models", "Data-driven chain rules (C1)", "Confident"),
    ("Mell, Scarfone, Romanosky - A Complete Guide to the Common Vulnerability Scoring System Version 2.0 (FIRST)", 2007, "Risk scoring", "CVSS base/temporal/environmental scoring", "-", "Severity bands and weighting", "Confident"),
    ("Wachter, Mittelstadt, Russell - Counterfactual Explanations without Opening the Black Box (Harvard JOLT)", 2017, "Counterfactual explanation", "Counterfactuals", "-", "What-If simulator framing (C2)", "Confident"),
    ("Wheeler - zxcvbn: Low-Budget Password Strength Estimation (USENIX Security)", 2016, "Password strength", "Pattern-matching entropy estimation", "Leaked password sets", "Pattern-penalised entropy", "Confident"),
    ("Weir, Aggarwal, Collins, Stern - Testing metrics for password creation policies by attacking large sets of revealed passwords (ACM CCS)", 2010, "Password strength", "Shows NIST entropy is a poor metric", "RockYou etc.", "Why naive entropy is not enough", "Confident"),
    ("Al-Subaiey et al. - Novel Interpretable and Robust Web-based AI Platform for Phishing Email Detection", 2024, "Phishing email detection", "Curated multi-corpus dataset + ML", "Nazario, SpamAssassin, CEAS, ...", "Email corpora mirror", "VERIFY author list / venue"),
]


def style(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def add_df(ws, df, start_row=None):
    if start_row is None:
        ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else (v.item() if hasattr(v, "item") else v) for v in r])


def main():
    wb = Workbook()
    # 1 Task tracker
    ws = wb.active
    ws.title = "Task Tracker"
    ws.append(["ID", "Module", "Task", "Owner", "Day", "Status", "Priority", "GitHub"])
    for i, (m, t, o, d, p, b) in enumerate(TASKS, 1):
        st = NOT_BUILT.get(t, BUILT_STATUS)
        ws.append([f"T{i:02d}", m, t, o, d, st, p, b])
        ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=STATUS_FILL[st])
    dv = DataValidation(type="list", formula1='"Not Started,In Progress,Blocked,Testing,Done"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{ws.max_row}")
    style(ws, [7, 16, 60, 11, 6, 13, 10, 20])

    # 2 Test cases (from real runs)
    ws = wb.create_sheet("Test Cases")
    tc = pd.read_csv(R / "test_cases.csv")
    tc = tc.rename(columns={"Pass/Fail": "Status"})[["Test ID", "Module", "Input (sanitized)", "Expected", "Actual", "Status", "Issue", "Fixed?"]]
    add_df(ws, tc)
    for row in ws.iter_rows(min_row=2):
        row[5].fill = PatternFill("solid", fgColor="C8E6C9" if row[5].value == "Pass" else "FFCDD2")
    style(ws, [9, 11, 48, 18, 16, 8, 40, 16])
    pt = R / "pytest_output.txt"
    if pt.exists():
        last = [ln for ln in pt.read_text().splitlines() if ln.strip()][-1]
        ws.append([])
        ws.append(["pytest", "all", "tests/ (unit + headless UI)", "all pass", last, "Pass" if "failed" not in last else "Fail"])

    # 3 ML results (real numbers only)
    ws = wb.create_sheet("ML Results")
    m = pd.read_csv(R / "ml_metrics.csv")
    m["Model"] = m.task.str.upper() + " - " + m.model + m.selected.map({True: " (selected)", False: ""})
    add_df(ws, m[["Model", "dataset", "accuracy", "precision", "recall", "f1", "n_train", "n_test", "tp", "fp", "fn", "tn"]]
           .rename(columns={"dataset": "Dataset", "accuracy": "Accuracy", "precision": "Precision", "recall": "Recall", "f1": "F1"}))
    ws.append([])
    ws.append(["Fusion selection (out-of-fold, training split only)"])
    add_df(ws, pd.read_csv(R / "fusion_selection.csv"))
    style(ws, [34, 60, 10, 10, 10, 10, 9, 9, 8, 8, 8, 8])

    # 4 Bugs
    ws = wb.create_sheet("Bugs")
    ws.append(["Bug ID", "Module", "Problem", "Severity", "Owner", "Status"])
    for b in BUGS:
        ws.append(list(b))
    style(ws, [9, 13, 70, 10, 11, 45])

    # 5 Literature
    ws = wb.create_sheet("Literature")
    ws.append(["Paper", "Year", "Problem", "Method", "Dataset", "What we use", "Confidence"])
    for row in LIT:
        ws.append(list(row))
    style(ws, [70, 7, 22, 36, 22, 36, 18])

    # 6 Novelty evidence
    ws = wb.create_sheet("Novelty Evidence")
    ws.append(["1. Scenario study (C1): plain weighted average vs Threat Correlation Engine"])
    ws["A1"].font = Font(bold=True)
    sc = pd.read_csv(R / "scenario_study.csv")
    sc["which_correct"] = sc.apply(lambda r: "both" if r.baseline_correct and r.correlated_correct else
                                   "correlated" if r.correlated_correct else "baseline" if r.baseline_correct else "neither", axis=1)
    add_df(ws, sc[["profile", "name", "baseline_score", "baseline_tier", "correlated_score", "correlated_tier", "expected",
                   "which_correct", "chains", "ablation_no_high_cap_tier"]])
    ws.append([(R / "scenario_study_summary.txt").read_text()])
    for title, f in [("2. Adversarial URL robustness (detection rate on 300 held-out phishing URLs)", "adversarial_url_results.csv"),
                     ("3. Rule-ML disagreement statistics (held-out test splits)", "disagreement_stats.csv"),
                     ("4. India-localized test set (28 hand-made messages: 21 scams, 7 legitimate)", "india_testset_results.csv")]:
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        add_df(ws, pd.read_csv(R / f))
    for i, w in enumerate([22, 46, 14, 14, 14, 14, 14, 16, 40, 18, 16, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
